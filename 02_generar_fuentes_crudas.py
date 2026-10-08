"""
02 · FUENTES CRUDAS "SUCIAS" (datos 100% ficticios)
Toma 200 pacientes de la cohorte limpia y los "desarma" en los archivos sueltos que una clínica real suele tener:
registro de pacientes, agenda, dos laboratorios (largo y ancho) + PDFs, Oligoscan mensual, D-ROMs/PIC cargados a mano,
ADN Test (solo PRO), formularios de anamnesis con texto libre y protocolos.
El pipeline (03) tiene que volver a ordenarlos; la cohorte limpia es la "verdad" para validarlo.
"""
import numpy as np, pandas as pd, random, os, json, unicodedata
from datetime import date, datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
COH = f'{HERE}/datos/cohorte'
OUT = f'{HERE}/datos/fuentes_crudas'
os.makedirs(f'{OUT}/pdfs_laboratorio', exist_ok=True)
rng = np.random.default_rng(99); random.seed(99)
noacc = lambda s: ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')

vis = pd.read_csv(f'{COH}/visitas_longitudinal.csv', parse_dates=['fecha_visita'])
pac = pd.read_csv(f'{COH}/pacientes.csv', parse_dates=['fecha_nacimiento'])
evt = pd.read_csv(f'{COH}/eventos_clinicos.csv')
SUBSET = sorted(rng.choice(pac.patient_id.values, 200, replace=False))
pac = pac[pac.patient_id.isin(SUBSET)].reset_index(drop=True)
vis = vis[vis.patient_id.isin(SUBSET)].reset_index(drop=True)
json.dump(SUBSET, open(f'{HERE}/datos/subset_ids.json', 'w'))
P = {r.patient_id: r for r in pac.itertuples()}
MEDICOS = ['Dr. Ledesma', 'Dra. Ibarra', 'Dr. Quiroga', 'Dra. Barrios']

# ------------------------------------------------------------------ suciedad
def name_variant(p, allow_initial=False):
    f, l = p.nombre, p.apellido
    c = rng.integers(0, 7)
    s = [f'{l.upper()}, {f}', f'{f} {l}', f'{l}, {f}', noacc(f'{f} {l}').lower(), noacc(f'{l} {f}').upper(), f'{f} {l} ', f'{l} {f}'][c]
    if allow_initial and rng.random() < 0.05: s = f'{f[0]}. {l}'
    if rng.random() < 0.03: s = s.replace('z', 's', 1) if 'z' in s else s.replace('ll', 'y', 1)
    return s
def dni_variant(dni):
    s = f'{int(dni):,}'.replace(',', '.'); c = rng.integers(0, 5)
    return [str(int(dni)), s, s.replace('.', ' '), int(dni), f'DNI {int(dni)}'][c]
def date_variant(d, styles=(0, 1, 2, 3)):
    d = pd.Timestamp(d); c = random.choice(styles)
    if c == 0: return datetime(d.year, d.month, d.day)
    if c == 1: return d.strftime('%d/%m/%Y')
    if c == 2: return d.strftime('%Y-%m-%d')
    return f'{d.day}/{d.month}/{d.year % 100:02d}'
def num_variant(x, p_comma=0.35):
    return str(x).replace('.', ',') if rng.random() < p_comma else x

ARIAL = Font(name='Arial', size=10); BOLD = Font(name='Arial', size=10, bold=True); HF = PatternFill('solid', fgColor='D9E1F2')
def put(ws, r, c, v, bold=False, fill=None):
    cell = ws.cell(row=r, column=c, value=v); cell.font = BOLD if bold else ARIAL
    if fill: cell.fill = fill
    if isinstance(v, datetime): cell.number_format = random.choice(['dd/mm/yyyy', 'yyyy-mm-dd', 'dd-mmm-yy'])
def widths(ws, ws_):
    for i, w in enumerate(ws_, 1): ws.column_dimensions[get_column_letter(i)].width = w
def header(ws, row, cols):
    for j, x in enumerate(cols, 1): put(ws, row, j, x, True, HF)

base = vis[vis.visit_n == 0].set_index('patient_id')
vis['fecha_estudio'] = vis.fecha_visita + pd.to_timedelta(rng.integers(0, 3, len(vis)), unit='D')
vis['has_oligo'] = rng.random(len(vis)) < np.where(vis.visit_n == 0, 0.95, 0.70)
vis['has_drompic'] = rng.random(len(vis)) < 0.92
def src(d):
    if d >= pd.Timestamp('2026-07-01') and rng.random() < 0.45: return 'PDF'
    if d >= pd.Timestamp('2026-04-01') and rng.random() < 0.40: return 'B'
    return 'A'
vis['lab_src'] = vis.fecha_estudio.map(src)
LAB = {'glucosa': ('mg/dL', '70 - 99', 0), 'hba1c': ('%', '4.0 - 5.6', 1), 'insulina': ('µU/mL', '2.6 - 24.9', 1), 'colesterol_total': ('mg/dL', '< 200', 0),
       'hdl': ('mg/dL', '> 40', 0), 'ldl': ('mg/dL', '< 130', 0), 'trigliceridos': ('mg/dL', '< 150', 0), 'pcr_us': ('mg/L', '< 3.0', 2),
       'homocisteina': ('µmol/L', '5.0 - 15.0', 1), 'vitamina_d': ('ng/mL', '30 - 100', 1), 'b12': ('pg/mL', '200 - 900', 0), 'tsh': ('µUI/mL', '0.4 - 4.0', 2),
       'ferritina': ('ng/mL', '15 - 300', 0), 'creatinina': ('mg/dL', '0.6 - 1.3', 2), 'alt': ('U/L', '< 41', 0), 'ast': ('U/L', '< 40', 0)}
NAMES_A = {'glucosa': ['Glucemia', 'GLUCOSA', 'Glucosa en ayunas'], 'hba1c': ['Hemoglobina glicosilada (HbA1c)', 'HbA1c', 'Hb glicosilada'], 'insulina': ['Insulina basal', 'INSULINA'],
           'colesterol_total': ['Colesterol total', 'COLESTEROL TOTAL'], 'hdl': ['HDL colesterol', 'Col HDL'], 'ldl': ['LDL colesterol', 'Col LDL (calculado)'],
           'trigliceridos': ['Triglicéridos', 'TRIGLICERIDOS'], 'pcr_us': ['PCR ultrasensible', 'Proteína C reactiva (us)', 'PCR-us'], 'homocisteina': ['Homocisteína', 'HOMOCISTEINA'],
           'vitamina_d': ['Vitamina D (25-OH)', '25 OH Vitamina D'], 'b12': ['Vitamina B12', 'B12'], 'tsh': ['TSH'], 'ferritina': ['Ferritina'], 'creatinina': ['Creatinina'],
           'alt': ['TGP (ALT)', 'ALT'], 'ast': ['TGO (AST)', 'AST']}
OLIGO = ['Calcio', 'Magnesio', 'Sodio', 'Potasio', 'Zinc', 'Cobre', 'Hierro', 'Selenio', 'Manganeso', 'Mercurio', 'Plomo', 'Cadmio', 'Aluminio', 'Arsénico', 'Níquel']
OSYM = dict(zip(OLIGO, ['Ca', 'Mg', 'Na', 'K', 'Zn', 'Cu', 'Fe', 'Se', 'Mn', 'Hg', 'Pb', 'Cd', 'Al', 'As', 'Ni']))
OREF = {'Calcio': (60, 120), 'Magnesio': (25, 50), 'Sodio': (700, 1000), 'Potasio': (250, 450), 'Zinc': (110, 200), 'Cobre': (10, 22), 'Hierro': (20, 45), 'Selenio': (0.9, 1.6),
        'Manganeso': (0.08, 0.2), 'Mercurio': (0, 0.8), 'Plomo': (0, 1.5), 'Cadmio': (0, 0.12), 'Aluminio': (0, 7.0), 'Arsénico': (0, 0.2), 'Níquel': (0, 0.9)}
okey = lambda k: 'oligo_' + noacc(k.lower())
# faltantes por visita (solo en las fuentes crudas)
miss_lab = {r.visit_id: {k for k in LAB if rng.random() < 0.03} for r in vis.itertuples()}
miss_ol = {r.visit_id: {k for k in OLIGO if rng.random() < 0.02} for r in vis.itertuples()}

# ================================================================ 01 registro de pacientes
wb = Workbook(); ws = wb.active; ws.title = 'Pacientes 2025'; ws2 = wb.create_sheet('Pacientes 2026')
header(ws, 1, ['DNI', 'Apellido', 'Nombre', 'Fecha de nacimiento', 'Sexo', 'Localidad', 'Programa', 'Fecha alta'])
header(ws2, 1, ['Documento', 'Apellido y nombre', 'Nac.', 'Género', 'Ciudad', 'Plan', 'Alta', 'Comentarios'])
SEXV = {'F': ['F', 'Femenino', 'mujer', 'fem'], 'M': ['M', 'Masculino', 'hombre', 'masc']}
PROG = {'BASE': ['BASE', 'Programa BASE', 'Base', 'base'], 'PRO': ['PRO', 'Programa PRO', 'Pro', 'pro']}
r1 = r2 = 2; dups = 0
for p in pac.itertuples():
    alta = base.loc[p.patient_id, 'fecha_visita'] - timedelta(days=int(rng.integers(0, 6)))
    sexo = random.choice(SEXV[p.sexo]) if rng.random() > 0.04 else None
    if alta.year == 2025:
        row = [dni_variant(p.dni), p.apellido if rng.random() > .2 else p.apellido.upper(), p.nombre, date_variant(p.fecha_nacimiento), sexo, p.localidad, random.choice(PROG[p.programa]), date_variant(alta, (0, 1))]
        for j, x in enumerate(row, 1): put(ws, r1, j, x)
        r1 += 1
        if rng.random() < 0.06:   # paciente cargado dos veces (otra persona de recepción, otro formato)
            row = [dni_variant(p.dni), p.apellido.upper(), noacc(p.nombre).lower(), date_variant(p.fecha_nacimiento), sexo, p.localidad, random.choice(PROG[p.programa]), date_variant(alta + timedelta(days=1), (0, 1))]
            for j, x in enumerate(row, 1): put(ws, r1, j, x)
            r1 += 1; dups += 1
    else:
        row = [dni_variant(p.dni), name_variant(p), date_variant(p.fecha_nacimiento), sexo, p.localidad, random.choice(PROG[p.programa]), date_variant(alta, (0, 1)), None]
        for j, x in enumerate(row, 1): put(ws2, r2, j, x)
        r2 += 1
        if rng.random() < 0.06:
            row = [dni_variant(p.dni), name_variant(p), date_variant(p.fecha_nacimiento), sexo, p.localidad, random.choice(PROG[p.programa]), date_variant(alta + timedelta(days=1), (0, 1)), 'posible duplicado?']
            for j, x in enumerate(row, 1): put(ws2, r2, j, x)
            r2 += 1; dups += 1
widths(ws, [14, 16, 14, 18, 11, 22, 12, 14]); widths(ws2, [14, 26, 14, 11, 22, 12, 14, 20])
wb.save(f'{OUT}/01_registro_pacientes.xlsx')

# ================================================================ 02 agenda y consultas
wb = Workbook(); ws = wb.active; ws.title = 'Agenda y consultas'
header(ws, 1, ['Fecha', 'DNI', 'Paciente', 'Tipo', 'Profesional', 'PA', 'Estado', 'Fecha estudios', 'Fecha integración', 'Notas'])
NOTAS = ['Refiere cansancio', 'Mejoró el sueño', 'Viene con estudios previos', 'Sin novedades', 'Consulta por energía y rendimiento', 'Pide control de seguimiento', 'Cambió trabajo, más estrés', None, None, None]
rr = 2; n_aus = 0
for v in vis.sort_values('fecha_visita').itertuples():
    p = P[v.patient_id]
    pa = f'{int(v.pa_sistolica)}/{int(76 + 0.5 * (v.pa_sistolica - 117) + rng.normal(0, 4))}'
    pa = random.choice([pa, pa.replace('/', ' / '), pa.replace('/', '-')]) if rng.random() > 0.07 else None
    fe = fi = None
    if v.visit_n == 0: fe, fi = date_variant(p.fecha_estudios, (0, 1)), date_variant(p.fecha_integracion, (0, 1))
    row = [date_variant(v.fecha_visita), dni_variant(p.dni), name_variant(p), 'Consulta integral' if v.visit_n == 0 else 'Seguimiento', random.choice(MEDICOS), pa, 'Asistió', fe, fi, random.choice(NOTAS)]
    for j, x in enumerate(row, 1): put(ws, rr, j, x)
    rr += 1
    if rng.random() < 0.05:
        d = v.fecha_visita + timedelta(days=int(rng.integers(30, 70)))
        row = [date_variant(d), dni_variant(p.dni), name_variant(p), 'Seguimiento', random.choice(MEDICOS), None, 'Ausente', None, None, 'No se presentó']
        for j, x in enumerate(row, 1): put(ws, rr, j, x)
        rr += 1; n_aus += 1
widths(ws, [14, 14, 26, 18, 14, 10, 10, 16, 18, 34]); wb.save(f'{OUT}/02_agenda_consultas.xlsx')

# ================================================================ 03 laboratorio
def lab_cell(kk, val):
    unit, ref, dec = LAB[kk]; u = unit
    if kk == 'glucosa' and rng.random() < 0.10: val, u, dec = val / 18.016, 'mmol/L', 2
    elif kk == 'pcr_us' and rng.random() < 0.15: val, u, dec = val / 10, 'mg/dL', 3
    elif kk == 'vitamina_d' and rng.random() < 0.08: val, u, dec = val * 2.496, 'nmol/L', 1
    if kk == 'pcr_us' and u == 'mg/L' and val < 0.3: return '<0,3', u
    res = num_variant(round(float(val), dec) if dec else int(round(val)))
    if rng.random() < 0.2: u = random.choice([u.lower(), u.upper()])
    return res, u

def build_A(year):
    wb = Workbook(); ws = wb.active; ws.title = 'Resultados'
    put(ws, 1, 1, 'LABORATORIO BIOQUÍMICO CENTRAL - Informe de resultados', True); ws.merge_cells('A1:G1')
    put(ws, 2, 1, f'Exportado el 01/10/2026  |  Período {year}')
    header(ws, 4, ['Fecha extracción', 'DNI', 'Paciente', 'Determinación', 'Resultado', 'Unidad', 'Valores de referencia'])
    rr, nd = 5, 0
    for v in vis[(vis.lab_src == 'A') & (vis.fecha_estudio.dt.year == year)].sort_values('fecha_estudio').itertuples():
        p = P[v.patient_id]; idd, idn, idf = dni_variant(p.dni), name_variant(p), date_variant(v.fecha_estudio, (0, 1, 3))
        for kk in LAB:
            if kk in miss_lab[v.visit_id]:
                if rng.random() < 0.5: continue
                res, u = random.choice(['Pendiente', 'N/R', None]), None
            else: res, u = lab_cell(kk, getattr(v, kk))
            row = [idf, idd, idn, random.choice(NAMES_A[kk]), res, u, LAB[kk][1]]
            for j, x in enumerate(row, 1): put(ws, rr, j, x)
            rr += 1; nd += 1
            if rng.random() < 0.01:
                for j, x in enumerate(row, 1): put(ws, rr, j, x)
                rr += 1
        if rng.random() < 0.15: rr += 1
    put(ws, rr + 1, 1, f'Total de determinaciones informadas: {nd}'); put(ws, rr + 2, 1, 'Bioq. M. Quiroga MP 12345 - Firma digital')
    widths(ws, [16, 14, 26, 32, 12, 10, 20]); wb.save(f'{OUT}/03_laboratorio_proveedorA_{year}.xlsx')
build_A(2025); build_A(2026)

wb = Workbook(); ws = wb.active; ws.title = 'Hoja1'
hB = ['N° Orden', 'Apellido y Nombre', 'DNI', 'F. Extracción', 'GLU (mg/dL)', 'HbA1c %', 'INS uU/mL', 'COL-T mg/dL', 'HDL mg/dL', 'LDL mg/dL', 'TG mg/dL', 'PCR-us mg/L',
      'HCY umol/L', '25-OH VIT D ng/mL', 'B12 pg/mL', 'TSH uUI/mL', 'FERR ng/mL', 'CREA mg/dL', 'TGP U/L', 'TGO U/L']
put(ws, 1, 1, 'Centro de Diagnóstico Norte - Exportación de órdenes', True); header(ws, 3, hB)
rr = 4
for v in vis[vis.lab_src == 'B'].sort_values('fecha_estudio').itertuples():
    p = P[v.patient_id]; row = [int(rng.integers(40000, 99999)), name_variant(p), dni_variant(p.dni), date_variant(v.fecha_estudio, (0, 1))]
    for kk in LAB: row.append(random.choice([None, 'NR', '-']) if kk in miss_lab[v.visit_id] else num_variant(getattr(v, kk), 0.5))
    for j, x in enumerate(row, 1): put(ws, rr, j, x)
    rr += 1
widths(ws, [10, 26, 14, 14] + [12] * 16); wb.save(f'{OUT}/03_laboratorio_proveedorB_2026.xlsx')

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
ss = getSampleStyleSheet(); npdf = 0
for v in vis[vis.lab_src == 'PDF'].itertuples():
    p = P[v.patient_id]
    fn = f"{OUT}/pdfs_laboratorio/informe_{noacc(p.apellido).upper()}_{v.fecha_estudio.strftime('%Y%m%d')}_{v.patient_id}.pdf"
    doc = SimpleDocTemplate(fn, pagesize=A4, topMargin=40, bottomMargin=40)
    el = [Paragraph('LABORATORIO BIOQUÍMICO CENTRAL', ss['Title']), Paragraph(f'Paciente: {p.apellido.upper()}, {p.nombre}', ss['Normal']),
          Paragraph(f'DNI: {int(p.dni):,}'.replace(',', '.') + f"  |  Sexo: {p.sexo}  |  Fecha de extracción: {v.fecha_estudio.strftime('%d/%m/%Y')}", ss['Normal']), Spacer(1, 14)]
    data = [['Determinación', 'Resultado', 'Unidad', 'Valores de referencia']]
    for kk in LAB:
        if kk in miss_lab[v.visit_id]: data.append([NAMES_A[kk][0], 'Pendiente', '', LAB[kk][1]])
        else:
            res, uu = lab_cell(kk, getattr(v, kk)); data.append([NAMES_A[kk][0], str(res).replace('.', ','), uu, LAB[kk][1]])
    t = Table(data, colWidths=[190, 80, 70, 130]); t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#D9E1F2')), ('GRID', (0, 0), (-1, -1), 0.3, colors.grey), ('FONTSIZE', (0, 0), (-1, -1), 9)]))
    doc.build(el + [t, Spacer(1, 16), Paragraph('Bioq. M. Quiroga MP 12345 - Documento sintético para fines de portfolio.', ss['Italic'])]); npdf += 1

# ================================================================ 04 oligoscan
wb = Workbook(); wb.remove(wb.active)
ov = vis[vis.has_oligo].copy(); ov['mes'] = ov.fecha_estudio.dt.strftime('%Y-%m')
for mth, g in sorted(ov.groupby('mes')):
    ws = wb.create_sheet(f'Oligoscan {mth}'); use_sym = rng.random() < 0.4; cols = OLIGO[:]
    if rng.random() < 0.5: random.shuffle(cols)
    put(ws, 1, 1, f'OLIGOSCAN - Exportación {mth} (valores y rangos ficticios, unidades arbitrarias)', True)
    header(ws, 3, ['Fecha', 'Paciente', 'DNI'] + [OSYM[c] if use_sym else c for c in cols]); put(ws, 4, 1, 'Ref. mín - máx')
    for j, c in enumerate(cols, 4): put(ws, 4, j, f'{OREF[c][0]} - {OREF[c][1]}')
    rr = 5
    for v in g.sort_values('fecha_estudio').itertuples():
        p = P[v.patient_id]; row = [date_variant(v.fecha_estudio, (0, 1)), name_variant(p), dni_variant(p.dni) if rng.random() > 0.35 else None]
        for c in cols: row.append('N/D' if c in miss_ol[v.visit_id] else num_variant(getattr(v, okey(c)), 0.25))
        for j, x in enumerate(row, 1): put(ws, rr, j, x)
        rr += 1
    widths(ws, [14, 26, 14] + [11] * len(cols))
wb.save(f'{OUT}/04_oligoscan_exportaciones.xlsx')

# ================================================================ 05 D-ROMs + PIC manual (sin DNI)
wb = Workbook(); ws = wb.active; ws.title = 'Registro 2025'; ws2 = wb.create_sheet('Registro 2026')
OBS = ['repetir', 'paciente ansioso', 'valor dudoso', 'ok', None, None, None, 'muestra tardía', 'ayuno 8h']
header(ws, 1, ['Fecha', 'Pac.', 'D-ROMs (U.CARR)', 'PIC', 'Obs']); header(ws2, 1, ['Día', 'Paciente', 'DROMs', 'PIC índice', 'Nivel PIC', 'Observaciones'])
r1 = r2 = 2
for v in vis[vis.has_drompic].sort_values('fecha_estudio').itertuples():
    p = P[v.patient_id]; nm = name_variant(p, True); dr = int(v.drom_ucarr) if rng.random() > .03 else None; pic = int(v.pic_indice) if rng.random() > .03 else None
    if v.fecha_estudio.year == 2025:
        row = [date_variant(v.fecha_estudio, (0, 1)), nm, dr, pic, random.choice(OBS)]
        for j, x in enumerate(row, 1): put(ws, r1, j, x)
        r1 += 1
    else:
        nivel = None if pic is None else ('bajo' if pic < 35 else 'moderado' if pic < 60 else 'alto')
        row = [date_variant(v.fecha_estudio, (0, 1, 3)), nm, dr, pic, nivel, random.choice(OBS)]
        for j, x in enumerate(row, 1): put(ws2, r2, j, x)
        r2 += 1
widths(ws, [14, 26, 16, 10, 20]); widths(ws2, [14, 26, 10, 12, 12, 20]); wb.save(f'{OUT}/05_drom_pic_registro_manual.xlsx')

# ================================================================ 06 ADN Test (solo PRO, primera visita)
def snp(g):
    return random.choice([g, g, '/'.join(g), g.lower() if g not in ('Ff', 'FF', 'ff') else g, ' '.join(g)]) if len(g) == 2 else g
wb = Workbook(); ws = wb.active; ws.title = 'ADN Test'
put(ws, 1, 1, 'ADN TEST - Longevidad & Metabolismo (datos 100% ficticios)', True)
header(ws, 3, ['Fecha informe', 'DNI', 'Paciente', 'APOE', 'MTHFR C677T', 'FTO rs9939609', 'ACE I/D', 'LCT -13910', 'HLA DQ2/DQ8', 'VDR FokI', 'CYP1A2 (cafeína)', 'GSTT1', 'Observaciones'])
rr = 4; FTOG = {0: 'TT', 1: 'AT', 2: 'AA'}
for v in vis[(vis.visit_n == 0) & (vis.programa == 'PRO')].itertuples():
    p = P[v.patient_id]
    ap = random.choice([v.apoe, v.apoe.upper(), v.apoe.replace('/', ''), v.apoe.replace('/', '-')])
    row = [date_variant(v.fecha_estudio + timedelta(days=int(rng.integers(10, 30))), (0, 1)), dni_variant(p.dni), name_variant(p), ap, snp(v.mthfr_c677t), snp(FTOG[int(v.fto_alelos_riesgo)]),
           snp(v.ace_id), snp(v.lct_13910), random.choice(['Positivo', '+', 'positivo']) if v.hla_dq2_dq8 == 1 else random.choice(['Negativo', '-', 'negativo']),
           random.choice([v.vdr_foki, v.vdr_foki.replace('F', 'F/', 1) if v.vdr_foki == 'Ff' else v.vdr_foki]), random.choice([v.cyp1a2_cafeina.capitalize(), v.cyp1a2_cafeina.upper(), v.cyp1a2_cafeina]),
           random.choice([v.gstt1.capitalize(), v.gstt1, 'nulo (deleción)' if v.gstt1 == 'nulo' else 'Presente']), None]
    for j, x in enumerate(row, 1): put(ws, rr, j, x)
    rr += 1
widths(ws, [14, 14, 26, 10, 14, 14, 10, 12, 12, 10, 14, 14, 14]); wb.save(f'{OUT}/06_adn_test_PRO.xlsx')

# ================================================================ 07 anamnesis (formularios con texto libre)
def hora_txt(h):
    mins = int(round(h * 60)) % (24 * 60); hh, mm = divmod(mins, 60)
    if rng.random() < 0.34: return f'{hh:02d}:{mm:02d}'
    if rng.random() < 0.5: return f'{(hh % 12) or 12}:{mm:02d} {"pm" if hh >= 12 else "am"}'
    return f'{hh}h{mm:02d}'
def like(x): s = f'{x:.1f}'; return random.choice([s, s.replace('.', ','), s + '/10', x]) if rng.random() > .04 else None
def ent(x, suf): return random.choice([int(x), str(int(x)), f'{int(x)} {suf}']) if rng.random() > .04 else None
FAMTXT = {'fam_cardiovascular': ['padre infarto', 'madre ACV', 'hermano hipertensión', 'abuelo cardiopatía'], 'fam_diabetes': ['madre diabetes', 'padre DBT2', 'abuela diabetes tipo 2'],
          'fam_cancer': ['abuelo cáncer de colon', 'madre cáncer de mama', 'tío tumor de próstata'], 'fam_neurodegenerativa': ['abuela Alzheimer', 'padre Parkinson', 'abuelo demencia']}
wb = Workbook(); ws = wb.active; ws.title = 'Formularios'
cols = ['Fecha', 'DNI', 'Paciente', 'Hora de acostarse', 'Celular en cama', 'Fuerza (sesiones/sem)', 'Cardio (min/sem)', 'Horas sentado/día', 'Ultraprocesados (comidas/sem)', 'Alcohol (tragos/sem)',
        'Dieta mediterránea (0-10)', 'Tabaco', 'Estrés (0-10)', 'Ansiedad (0-10)', 'Calidad de sueño (0-10)', 'Energía (0-10)', 'Dolor (0-10)', 'Claridad mental (0-10)', 'Molestias digestivas (0-10)',
        'Antecedentes familiares', 'Intolerancias', 'Cirugías / internaciones / diagnósticos previos']
header(ws, 1, cols); rr = 2
for v in vis.sort_values('fecha_visita').itertuples():
    p = P[v.patient_id]
    cel = random.choice(['Sí', 'si', 'SI', 'a veces']) if v.celular_en_cama == 1 else random.choice(['No', 'no', 'NO', 'nunca'])
    tab = {'Nunca': random.choice(['Nunca', 'No fuma']), 'Ex': random.choice(['Ex fumador', 'ex']), 'Actual': random.choice(['Fuma', 'Sí, actualmente'])}[v.tabaquismo]
    fam = intol = ev = None
    if v.visit_n == 0:
        items = [random.choice(FAMTXT[k]) for k in FAMTXT if getattr(p, k) == 1]
        fam = random.choice(['Sin antecedentes', 'Ninguno relevante']) if not items else random.choice(['; ', ' / ']).join(items)
        it = []
        if p.intol_gluten_detectada: it.append(random.choice(['sensibilidad al gluten', 'no tolera el gluten (trigo)', 'intolerancia al gluten']))
        if p.intol_lactosa_detectada: it.append(random.choice(['no tolera lactosa', 'intolerancia a la lactosa', 'mala digestión de lácteos']))
        intol = '; '.join(it) if it else random.choice(['Ninguna', 'No refiere', '-', None])
        es = evt[evt.patient_id == v.patient_id]
        ev = random.choice(['Sin antecedentes', 'No refiere']) if es.empty else random.choice(['; ', ' / ']).join(random.choice([f'{e.tipo_evento} {e.anio}', f'{e.tipo_evento.capitalize()} ({e.anio})', f'{noacc(e.tipo_evento)} - {e.anio}']) for e in es.itertuples())
    row = [date_variant(v.fecha_visita), dni_variant(p.dni), name_variant(p), hora_txt(v.hora_acostarse), cel, ent(v.sesiones_fuerza_sem, 'veces'), ent(v.min_cardio_sem, 'min'),
           random.choice([num_variant(v.sedentarismo_h_dia, .4), f'{str(v.sedentarismo_h_dia).replace(".", ",")} hs']), ent(v.ultraprocesados_sem, 'veces'), ent(v.alcohol_sem, 'tragos'),
           like(v.dieta_mediterranea_score), tab, like(v.estres_percibido), like(v.ansiedad_score), like(v.calidad_sueno), like(v.energia), like(v.dolor), like(v.claridad_mental), like(v.molestias_digestivas), fam, intol, ev]
    for j, x in enumerate(row, 1): put(ws, rr, j, x)
    rr += 1
widths(ws, [14, 14, 26, 14, 12, 14, 14, 14, 16, 14, 14, 14, 10, 10, 12, 10, 10, 12, 14, 40, 34, 60]); wb.save(f'{OUT}/07_anamnesis_formularios.xlsx')

# ================================================================ 08 protocolos
wb = Workbook(); ws = wb.active; ws.title = 'Protocolos'
header(ws, 1, ['DNI', 'Paciente', 'Fecha inicio', 'Protocolo', 'Adherencia reportada', 'Hallazgo (consulta de integración)'])
ADHV = {'Alta': ['Alta', 'alta', 'Buena', '90%'], 'Media': ['Media', 'media', 'Regular', '60%'], 'Baja': ['Baja', 'baja', 'Mala', '20%']}
rr = 2
for p in pac.itertuples():
    b = base.loc[p.patient_id]; d0 = b.fecha_visita + timedelta(days=int(rng.integers(0, 8)))
    hal = b.hallazgo_principal
    row = [dni_variant(p.dni), name_variant(p), date_variant(d0, (0, 1, 2)), random.choice([p.protocolo, p.protocolo.upper(), p.protocolo.lower()]),
           random.choice(ADHV[p.adherencia]) if rng.random() > .15 else None, None if hal == 'ninguno' else random.choice([hal, hal.capitalize(), hal.upper()])]
    for j, x in enumerate(row, 1): put(ws, rr, j, x)
    rr += 1
widths(ws, [14, 26, 14, 30, 18, 36]); wb.save(f'{OUT}/08_protocolos.xlsx')
print('pacientes', len(pac), '| visitas', len(vis), '| fuentes lab:', vis.lab_src.value_counts().to_dict(), '| PDFs', npdf, '| duplicados registro', dups, '| ausentes', n_aus)
